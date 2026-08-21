"""Independent verification of Bloomfield_Truck_Parking_Model.xlsx.

LibreOffice is unusable in this container - it fails even on a four-formula
workbook - so the xlsx skill's recalc.py cannot be used here. This script
substitutes for it: it evaluates every formula with the pure-Python `formulas`
engine, reports any cell that resolves to an Excel error, and diffs the
Summary tab against verify.py, which reimplements the same model from scratch.

Run:  python3 verify_workbook.py      (exit 0 = pass, 1 = fail)
"""
import contextlib, io, re, sys, warnings

warnings.filterwarnings("ignore")
sys.path.insert(0, ".")

import formulas
from openpyxl import load_workbook

BOOK = "Bloomfield_Truck_Parking_Model.xlsx"

_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    import verify as V

# ---------------------------------------------------------------- evaluate
sol = formulas.ExcelModel().loads(BOOK).finish().calculate()

# Keys look like "'[Bloomfield_Truck_Parking_Model.xlsx]SUMMARY'!C12".
# The book name keeps its original case; sheet names come back uppercased.
KEY = re.compile(r"^'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$")

sheets, errors = {}, {}
for key, node in sol.items():
    m = KEY.match(key)
    if not m:
        continue
    try:
        val = node.value[0, 0]
    except Exception:
        continue
    sheets.setdefault(m.group(1).upper(), {})[m.group(2)] = val
    if str(val).startswith("#"):
        errors.setdefault(str(val), []).append(f"{m.group(1)}!{m.group(2)}")

n_cells = sum(len(v) for v in sheets.values())
n_errors = sum(len(v) for v in errors.values())
print(f"cells evaluated : {n_cells:,}")
print(f"formula errors  : {n_errors or 'NONE'}")
for kind, where in errors.items():
    print("   ", kind, where[:15], "..." if len(where) > 15 else "")

# ------------------------------------------------- diff Summary vs verify.py
# Summary row labels -> verify.py result keys. Columns C/D/E are S1/S2/S3.
ROWS = {
    "truck stalls": ("stalls", 0),
    "fixed capex": ("alloc_f", 0),
    "variable capex": ("alloc_v", 0),
    "total project cost": ("total", 0),
    "fixed cost per stall": ("cps_f", 0),
    "variable cost per stall": ("cps_v", 0),
    "all-in cost per stall": ("cps_t", 0),
    "effective gross income": ("egi", 0),
    "operating expenses": ("opex", 0),
    "net operating income": ("noi", 0),
    "yield on cost": ("yoc", 6),
    "stabilized value at exit cap": ("val", 0),
    "development profit / (loss)": ("profit", 0),
    "loan at target 65% ltc": ("l_ltc", 0),
    "annual debt service": ("ds65", 0),
    "dscr at target leverage": ("dscr65", 6),
    "debt yield": ("dy", 6),
    "supportable loan (lesser of ltc / ltv / dscr)": ("loan", 0),
    "required equity": ("eq", 0),
    "required blended revenue per stall": ("req_rev", 2),
    "maximum supportable cost per stall": ("max_cps", 0),
    "maximum supportable land price": ("max_land", 0),
    "break-even occupancy": ("be_occ", 6),
}

ws = load_workbook(BOOK)["Summary"]
label_rows = {}
for r in range(1, ws.max_row + 1):
    lab = ws.cell(row=r, column=1).value
    if isinstance(lab, str):
        label_rows.setdefault(lab.strip().lower(), r)

summary = sheets.get("SUMMARY", {})
mismatches, missing, compared = [], [], 0
for label, (key, places) in ROWS.items():
    row = label_rows.get(label)
    if row is None:
        missing.append(f"label not found: {label!r}")
        continue
    for i, scen in enumerate(("S1", "S2", "S3")):
        got = summary.get(f"{'CDE'[i]}{row}")
        if got is None:
            missing.append(f"{label} [{scen}] not evaluated")
            continue
        want = V.R[scen][key]
        compared += 1
        tol = 10 ** -places if places else 0.51
        if abs(float(got) - float(want)) > tol:
            mismatches.append(f"{label} [{scen}]  workbook={got!r}  model={want!r}")

print(f"\nSummary cells compared against verify.py: {compared}"
      f"  |  mismatches: {len(mismatches)}")
for m in mismatches:
    print("   MISMATCH:", m)
for m in missing:
    print("   MISSING :", m)

expected = len(ROWS) * 3
ok = (n_errors == 0 and not mismatches and not missing and compared == expected)
if compared != expected:
    print(f"\n!! only {compared} of {expected} expected comparisons ran")
print("\nRESULT:", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
